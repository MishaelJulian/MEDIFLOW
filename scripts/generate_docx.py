import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def format_inline_runs(paragraph, text):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif token.startswith('[') and ']' in token and '(' in token and token.endswith(')'):
            match = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if match:
                link_text = match.group(1)
                run = paragraph.add_run(link_text)
                run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
                run.underline = True
            else:
                paragraph.add_run(token)
        else:
            paragraph.add_run(token)

def build_docx(md_path, docx_path):
    doc = Document()
    base_dir = os.path.dirname(os.path.abspath(md_path))

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        
        parsed_rows = []
        for line in table_lines:
            line = line.strip()
            if not line or line.startswith('|---') or line.startswith('|:---') or line.startswith('|-'):
                continue
            cells = [c.strip() for c in line.split('|')[1:-1]]
            if cells:
                parsed_rows.append(cells)
        
        if parsed_rows:
            num_cols = max(len(r) for r in parsed_rows)
            table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            add_table_borders(table)

            for row_idx, row_data in enumerate(parsed_rows):
                row = table.rows[row_idx]
                make_row_cant_split(row)
                is_header = (row_idx == 0)
                
                for col_idx in range(num_cols):
                    cell = row.cells[col_idx]
                    cell_text = row_data[col_idx] if col_idx < len(row_data) else ""
                    set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                    
                    if is_header:
                        set_cell_background(cell, "F1F5F9")
                    else:
                        if row_idx % 2 == 1:
                            set_cell_background(cell, "FAFAFA")

                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.15
                    format_inline_runs(p, cell_text)
                    if is_header:
                        for run in p.runs:
                            run.bold = True
                            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_before = Pt(4)
            p_after.paragraph_format.space_after = Pt(6)

        table_lines = []

    def flush_code_block():
        nonlocal code_lines
        if not code_lines:
            return
        
        code_text = "".join(code_lines).rstrip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.1
        
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        run = cp.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(2)
        p_after.paragraph_format.space_after = Pt(4)
        
        code_lines = []

    for line in lines:
        stripped = line.strip()

        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                flush_code_block()
            else:
                if in_table:
                    in_table = False
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            in_table = False
            flush_table()

        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', stripped)
        if img_match:
            caption = img_match.group(1)
            img_rel_path = img_match.group(2)
            img_full_path = os.path.normpath(os.path.join(base_dir, img_rel_path))

            if os.path.exists(img_full_path):
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.paragraph_format.space_before = Pt(10)
                img_p.paragraph_format.space_after = Pt(2)
                
                run = img_p.add_run()
                run.add_picture(img_full_path, width=Inches(5.8))

                if caption:
                    cap_p = doc.add_paragraph()
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_p.paragraph_format.space_before = Pt(2)
                    cap_p.paragraph_format.space_after = Pt(10)
                    cap_run = cap_p.add_run(caption)
                    cap_run.font.size = Pt(9.5)
                    cap_run.font.italic = True
                    cap_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
            continue

        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            continue

        if not stripped:
            continue

        if stripped.startswith('# '):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            format_inline_runs(h, stripped[2:])
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(17)
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                r.bold = True
        elif stripped.startswith('## '):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            format_inline_runs(h, stripped[3:])
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(13.5)
                r.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
                r.bold = True
        elif stripped.startswith('### '):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
            format_inline_runs(h, stripped[4:])
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(11.5)
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                r.bold = True
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.15
            format_inline_runs(p, stripped[2:])
        elif re.match(r'^\d+\.\s', stripped):
            match = re.match(r'^(\d+\.)\s(.*)', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.15
            format_inline_runs(p, match.group(2))
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2.5)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            format_inline_runs(p, stripped)

    if in_table:
        flush_table()
    if in_code_block:
        flush_code_block()

    saved_paths = []
    for target in [docx_path, docx_path.replace('.docx', '_FINAL.docx'), docx_path.replace('.docx', '_ADSE.docx')]:
        try:
            doc.save(target)
            saved_paths.append(target)
        except PermissionError:
            pass
            
    print(f"Saved to: {saved_paths}")

if __name__ == '__main__':
    md_file = os.path.join(os.path.dirname(__file__), '..', 'REPORT.md')
    docx_file = os.path.join(os.path.dirname(__file__), '..', 'P02_MEDIFLOW_REPORT.docx')
    build_docx(md_file, docx_file)
